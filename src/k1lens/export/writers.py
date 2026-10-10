"""CSV and Excel exports. Money is written as decimal strings so cents survive."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from k1lens.domain.access import coded_amounts, get_extracted
from k1lens.domain.fields import CODED_BOXES, SCALAR_FIELDS, FieldKind
from k1lens.domain.models import ScheduleK1
from k1lens.domain.parsing import format_value
from k1lens.validate.engine import Finding


@dataclass(frozen=True, slots=True)
class ExportRecord:
    document: ScheduleK1
    approved_by: str
    approved_at: str
    findings: tuple[Finding, ...] = ()


def headers() -> list[str]:
    columns = ["document_id", "tax_year", "variant"]
    columns.extend(spec.pdf_label for spec in SCALAR_FIELDS)
    columns.extend(spec.pdf_prefix for spec in CODED_BOXES)
    columns.extend(["approved_by", "approved_at", "review_findings"])
    return columns


_PLAIN_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")


def excel_safe(value: str) -> str:
    """Stop a cell that starts with a formula character from being evaluated."""

    if value == "":
        return value
    if value[0] in "=+@\t\r":
        return "'" + value
    if value.startswith("-") and _PLAIN_NUMBER.fullmatch(value) is None:
        return "'" + value
    return value


def row_values(record: ExportRecord) -> list[str]:
    document = record.document
    values = [
        document.document_id,
        "" if document.tax_year is None else str(document.tax_year),
        document.variant,
    ]
    for spec in SCALAR_FIELDS:
        values.append(format_value(spec.kind, get_extracted(document, spec.path).value))
    for box in CODED_BOXES:
        parts = [
            f"{code}={format_value(FieldKind.MONEY, amount)}"
            for code, amount in coded_amounts(document, box.path)
            if amount is not None
        ]
        values.append("; ".join(parts))
    values.append(record.approved_by)
    values.append(record.approved_at)
    values.append("; ".join(f"{item.rule_id} {item.field_path}" for item in record.findings))
    return [excel_safe(value) for value in values]


def records_to_csv(records: list[ExportRecord]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(headers())
    for record in records:
        writer.writerow(row_values(record))
    return buffer.getvalue()


def records_to_xlsx(records: list[ExportRecord]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        sheet = workbook.create_sheet("approved_k1s")
    sheet.title = "approved_k1s"
    header = headers()
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        sheet.append(row_values(record))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, name in enumerate(header, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(name) + 2, 14), 36)
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()
