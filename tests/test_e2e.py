"""generate -> extract -> validate -> review -> export."""

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from k1lens.domain.models import ScheduleK1
from k1lens.extract.deterministic import DeterministicExtractor
from k1lens.generate.build import Variant
from k1lens.generate.corpus import generate_corpus, load_ground_truth
from k1lens.review.app import create_app
from k1lens.roster import dump_roster, roster_from_documents
from k1lens.validate.engine import validate_document


def test_pipeline(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    generate_corpus(
        corpus,
        seed=42,
        count=2,
        plan=(Variant.CLEAN, Variant.INCONSISTENT_CAPITAL),
    )
    extractor = DeterministicExtractor()
    clean_pdf = corpus / "K1-0001.pdf"
    extracted = extractor.extract(clean_pdf)
    truth = load_ground_truth(corpus / "K1-0001.json").document
    assert extracted.part_i.ein.value == truth.part_i.ein.value
    bare_errors = [
        item.rule_id for item in validate_document(extracted) if item.severity.value == "error"
    ]
    assert set(bare_errors) == {"no_roster"}
    inconsistent = extractor.extract(corpus / "K1-0002.pdf")
    assert any(item.rule_id == "capital_reconciles" for item in validate_document(inconsistent))

    application = create_app(db_path=tmp_path / "review.db", fixture_dir=corpus, auto_ingest=True)
    with TestClient(application) as client:
        assert "K1-0002" in client.get("/").text
        saved = client.post(
            "/api/documents/K1-0001/corrections",
            json={"field_path": "part_i.city", "value": "Nashville", "actor": "Ada Reviewer"},
        )
        assert saved.status_code == 200
        # A person can still replace the generated partnership name.
        renamed = client.post(
            "/api/documents/K1-0001/corrections",
            json={
                "field_path": "part_i.name",
                "value": "SAMPLE Redcedar Ventures LP",
                "actor": "Ada Reviewer",
            },
        )
        assert renamed.status_code == 200
        partner = client.get("/api/documents/K1-0001").json()["document"]["part_ii"]["name"][
            "value"
        ]
        confirmed = client.post(
            "/api/documents/K1-0001/corrections",
            json={"field_path": "part_ii.name", "value": partner, "actor": "Ada Reviewer"},
        )
        assert confirmed.status_code == 200
        current = client.get("/api/documents/K1-0001").json()["document"]
        roster = roster_from_documents([ScheduleK1.model_validate(current)])
        assert client.post("/api/roster", json=dump_roster(roster)).status_code == 200
        approved = client.post(
            "/documents/K1-0001/approve", data={"actor": "Ada Reviewer"}, follow_redirects=True
        )
        assert approved.status_code == 200
        assert "approved" in approved.text
        audit = client.get("/api/documents/K1-0001/audit").json()
        assert any(
            entry["action"] == "approve" and entry["actor"] == "Ada Reviewer" for entry in audit
        )
        csv_export = client.get("/api/export/csv")
        assert "K1-0001" in csv_export.text
        assert "Nashville" in csv_export.text
        assert "K1-0002" not in csv_export.text
        sheet = load_workbook(BytesIO(client.get("/api/export/xlsx").content))["approved_k1s"]
        headers = [cell.value for cell in sheet[1]]
        assert "B Partnership city" in headers
        city_column = headers.index("B Partnership city") + 1
        assert sheet.cell(2, city_column).value == "Nashville"
