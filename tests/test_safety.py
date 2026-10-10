"""Regression tests: amounts, duplicates, review, and export."""

import random
import sqlite3
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook
from PIL import Image
from reportlab.pdfgen.canvas import Canvas

from k1lens.domain.access import apply_correction, coded_amounts, get_coded, get_extracted
from k1lens.domain.fields import REVIEW_THRESHOLD
from k1lens.domain.models import ScheduleK1
from k1lens.domain.parsing import parse_decimal
from k1lens.export.writers import ExportRecord, excel_safe, records_to_csv, records_to_xlsx
from k1lens.extract.assemble import assemble
from k1lens.extract.confidence import below_review_threshold
from k1lens.extract.deterministic import DeterministicExtractor
from k1lens.extract.text import Observed
from k1lens.generate.build import Variant, build_document
from k1lens.review.app import create_app
from k1lens.review.errors import ApprovalBlockedError, InvalidFieldValueError
from k1lens.review.preview import pdf_preview_png
from k1lens.review.store import ReviewStore, require_actor, verify_database
from k1lens.validate.engine import validate_document

MONEY = st.decimals(
    min_value="0", max_value="999999", places=2, allow_nan=False, allow_infinity=False
)


def _write_pdf(path: Path, lines: list[str]) -> None:
    canvas = Canvas(str(path))
    y = 740
    for line in lines:
        canvas.drawString(72, y, line)
        y -= 16
    canvas.save()


def test_parse_decimal_keeps_a_footnote_and_rejects_two_amounts() -> None:
    assert parse_decimal("1,234*") == Decimal("1234")
    assert parse_decimal("1,234 *") == Decimal("1234")
    assert parse_decimal("1,234.50(1)") == Decimal("1234.50")
    assert parse_decimal("(15.25)") == Decimal("-15.25")
    assert parse_decimal("$1,234.56") == Decimal("1234.56")
    with pytest.raises(ValueError, match="multiple"):
        parse_decimal("1,000 2,000")
    with pytest.raises(ValueError, match="multiple"):
        parse_decimal("12,345 1")
    with pytest.raises(ValueError):
        parse_decimal("NaN")
    with pytest.raises(ValueError):
        parse_decimal("Infinity")
    with pytest.raises(ValueError, match="US number"):
        parse_decimal("12.345,67")
    with pytest.raises(ValueError, match="scientific"):
        parse_decimal("1E3")
    with pytest.raises(ValueError, match="scientific"):
        parse_decimal("1e3")
    with pytest.raises(ValueError, match="thousands"):
        parse_decimal("0,50")
    with pytest.raises(ValueError, match="thousands"):
        parse_decimal("12,34")
    with pytest.raises(ValueError, match="thousands"):
        parse_decimal("1,2,3")
    assert parse_decimal(",17761.00") == Decimal("17761.00")
    assert parse_decimal("17761.00,") == Decimal("17761.00")


@given(left=MONEY, right=MONEY)
@settings(max_examples=40, deadline=None)
def test_two_amounts_never_concatenate(left: Decimal, right: Decimal) -> None:
    raw = f"{left:,.2f} {right:,.2f}"
    with pytest.raises(ValueError):
        parse_decimal(raw)


def test_nan_on_a_page_does_not_crash(tmp_path: Path) -> None:
    pdf = tmp_path / "nan.pdf"
    _write_pdf(
        pdf,
        [
            "A Partnership EIN: 00-0000001",
            "B Partnership name: SAMPLE One LP",
            "1 Ordinary business income: 10.00",
            "5 Interest income: NaN",
        ],
    )
    document = DeterministicExtractor().extract(pdf)
    assert document.part_iii.interest_income.value is None


def test_duplicate_labels_fall_below_the_threshold(tmp_path: Path) -> None:
    pdf = tmp_path / "dup.pdf"
    _write_pdf(
        pdf,
        [
            "A Partnership EIN: 00-0000001",
            "B Partnership name: SAMPLE One LP",
            "1 Ordinary business income: 10.00",
            "5 Interest income: 1,000",
            "5 Interest income: 2,000",
        ],
    )
    document = DeterministicExtractor().extract(pdf)
    interest = document.part_iii.interest_income
    assert interest.confidence < REVIEW_THRESHOLD
    assert interest.value in {None, Decimal("1000"), Decimal("2000"), "1000", "2000"} or str(
        interest.value
    ) in {"1000", "2000", "1000.00", "2000.00"}


def test_duplicate_coded_lines_are_both_kept(tmp_path: Path) -> None:
    pdf = tmp_path / "coded.pdf"
    _write_pdf(
        pdf,
        [
            "A Partnership EIN: 00-0000001",
            "B Partnership name: SAMPLE One LP",
            "1 Ordinary business income: 10.00",
            "13 Other deductions W: 100.00",
            "13 Other deductions W: 200.00",
        ],
    )
    document = DeterministicExtractor().extract(pdf)
    amounts = coded_amounts(document, "part_iii.other_deductions")
    assert amounts == [("W", Decimal("100.00")), ("W", Decimal("200.00"))] or _decimals(amounts)


def _decimals(amounts: list[tuple[str, Decimal | None]]) -> bool:
    assert [code for code, _value in amounts] == ["W", "W"]
    assert [Decimal(str(value)) for _code, value in amounts] == [
        Decimal("100.00"),
        Decimal("200.00"),
    ]
    return True


def test_coded_model_keeps_two_entries_with_the_same_code() -> None:
    document = build_document(random.Random(3), 1, Variant.CLEAN)
    payload = document.model_dump(mode="json")
    blank = {
        "value": "10.00",
        "confidence": 1.0,
        "provenance": {"page": 1, "bbox": None, "field_name": "13", "extractor": "test"},
    }
    payload["part_iii"]["other_deductions"] = [
        {"code": "W", "description": None, "amount": blank},
        {"code": "W", "description": None, "amount": {**blank, "value": "20.00"}},
    ]
    loaded = ScheduleK1.model_validate(payload)
    assert coded_amounts(loaded, "part_iii.other_deductions") == [
        ("W", Decimal("10.00")),
        ("W", Decimal("20.00")),
    ]
    updated = apply_correction(loaded, "part_iii.other_deductions[W]#2", "30", "Ada Reviewer")
    assert get_coded(updated, "part_iii.other_deductions")[1].amount.value == Decimal("30")


def test_reingest_keeps_the_corrected_record_and_approval(tmp_path: Path) -> None:
    from k1lens.roster import roster_from_documents

    store = ReviewStore(tmp_path / "review.db")
    document = build_document(random.Random(8), 1, Variant.CLEAN)
    pdf = b"%PDF-1.4 synthetic"
    store.upsert(document, pdf, [])
    store.correct(document.document_id, "part_i.city", "Nashville", "Ada Reviewer")
    corrected = store.get(document.document_id)
    assert corrected is not None
    store.set_roster(roster_from_documents([corrected.document]))
    store.approve(document.document_id, "Ada Reviewer")
    fresh = document.model_copy(deep=True)
    store.upsert(fresh, pdf, [])
    stored = store.get(document.document_id)
    assert stored is not None
    assert stored.status == "approved"
    assert stored.approved_by == "Ada Reviewer"
    assert stored.document.part_i.city.value == "Nashville"
    city = get_extracted(stored.document, "part_i.city")
    assert city.provenance.extractor == "human:Ada Reviewer"
    actions = [entry.action for entry in store.audit(document.document_id)]
    assert actions.count("ingest") == 2
    assert actions.count("correct") == 1
    assert actions.count("approve") == 1


def test_reingest_of_a_changed_approved_record_returns_it_to_review(tmp_path: Path) -> None:
    """A new amount on an approved record drops the approval and is audited.

    ``pending`` is the needs-review status. The audit row stores the old and
    new values. A human correction on another field stays on the record.
    """

    from k1lens.roster import roster_from_documents

    store = ReviewStore(tmp_path / "changed.db")
    document = build_document(random.Random(8), 1, Variant.CLEAN)
    pdf = b"%PDF-1.4 synthetic"
    store.upsert(document, pdf, [])
    store.correct(document.document_id, "part_i.city", "Nashville", "Ada Reviewer")
    corrected = store.get(document.document_id)
    assert corrected is not None
    store.set_roster(roster_from_documents([corrected.document]))
    store.approve(document.document_id, "Ada Reviewer")
    original = get_extracted(document, "part_iii.ordinary_business_income").value
    assert original is not None
    payload = document.model_dump(mode="json")
    payload["part_iii"]["ordinary_business_income"]["value"] = "999999.00"
    changed = ScheduleK1.model_validate(payload)
    store.upsert(changed, pdf, [])
    stored = store.get(document.document_id)
    assert stored is not None
    assert stored.status == "pending"
    assert stored.approved_by is None
    assert stored.approved_at is None
    assert stored.document.part_i.city.value == "Nashville"
    assert stored.document.part_iii.ordinary_business_income.value == Decimal("999999.00")
    row = next(
        entry
        for entry in store.audit(document.document_id)
        if entry.field_path == "part_iii.ordinary_business_income"
    )
    assert row.old_value is not None and str(original) in row.old_value.replace(",", "")
    assert row.new_value == "999999.00"


def test_approve_and_export_wait_for_an_override(tmp_path: Path) -> None:
    from k1lens.roster import roster_from_documents

    document = build_document(random.Random(4), 2, Variant.CLEAN)
    store = ReviewStore(tmp_path / "gate.db", roster=roster_from_documents([document]))
    payload = document.model_dump(mode="json")
    payload["part_iii"]["interest_income"]["confidence"] = 0.2
    low = ScheduleK1.model_validate(payload)
    store.upsert(low, b"%PDF-1.4", [])
    with pytest.raises(ApprovalBlockedError):
        store.approve(low.document_id, "Ada Reviewer")
    store.override(
        low.document_id,
        "part_iii.interest_income",
        "Checked the statement",
        "Ada Reviewer",
    )
    approved = store.approve(low.document_id, "Ada Reviewer")
    assert approved.status == "approved"
    snapshot = next(entry for entry in store.audit(low.document_id) if entry.action == "approve")
    assert "interest_income" in (snapshot.new_value or "")
    assert '"status": "approved"' in (snapshot.new_value or "")
    store.assert_exportable()


def test_an_ocr_page_failure_blocks_approval_until_a_person_overrides(tmp_path: Path) -> None:
    """A confident record still cannot be approved when a page failed OCR.

    Overriding the page list is the explicit review. A later ingest that
    fails a different set of pages drops that override and the approval.
    """

    document = build_document(random.Random(4), 2, Variant.CLEAN).model_copy(
        update={"ocr_failed_pages": [2]}
    )
    from k1lens.roster import roster_from_documents

    covering = roster_from_documents([document])
    store = ReviewStore(tmp_path / "ocr.db", roster=covering)
    findings = validate_document(document, roster=covering)
    assert [item.rule_id for item in findings if item.severity.value == "error"] == ["ocr_page"]
    store.upsert(document, b"%PDF-1.4", findings)
    application = create_app(db_path=store.db_path, auto_ingest=False)
    with TestClient(application) as client:
        page = client.get(f"/documents/{document.document_id}")
        assert page.status_code == 200
        assert "OCR failed on page 2" in page.text
        assert "ocr_page" in page.text
        blocked = client.post(
            f"/documents/{document.document_id}/approve",
            data={"actor": "Ada Reviewer"},
        )
        assert blocked.status_code == 400
        assert "OCR failed on page 2" in blocked.text
    with pytest.raises(ApprovalBlockedError, match="OCR failed on page 2"):
        store.approve(document.document_id, "Ada Reviewer")
    corrected = store.correct(document.document_id, "part_i.city", "Nashville", "Ada Reviewer")
    assert corrected.document.ocr_failed_pages == [2]
    assert any(item.rule_id == "ocr_page" for item in corrected.findings)
    store.set_roster(roster_from_documents([corrected.document]))
    with pytest.raises(InvalidFieldValueError, match="overridden, not edited"):
        store.correct(document.document_id, "ocr_failed_pages", "2", "Ada Reviewer")
    store.override(
        document.document_id,
        "ocr_failed_pages",
        "Read page 2 from the scan",
        "Ada Reviewer",
    )
    approved = store.approve(document.document_id, "Ada Reviewer")
    assert approved.status == "approved"
    again = document.model_copy(update={"ocr_failed_pages": [2, 3]})
    store.upsert(again, b"%PDF-1.4", validate_document(again))
    stored = store.get(document.document_id)
    assert stored is not None
    assert stored.status == "pending"
    assert stored.document.ocr_failed_pages == [2, 3]
    assert stored.document.part_i.city.value == "Nashville"
    with pytest.raises(ApprovalBlockedError, match=r"OCR failed on pages 2, 3"):
        store.approve(document.document_id, "Ada Reviewer")


def test_literal_reviewer_is_rejected() -> None:
    with pytest.raises(InvalidFieldValueError):
        require_actor("reviewer")
    with pytest.raises(InvalidFieldValueError):
        require_actor("  ")


def test_audit_log_is_append_only_and_chained(tmp_path: Path) -> None:
    db = tmp_path / "audit.db"
    store = ReviewStore(db)
    document = build_document(random.Random(5), 1, Variant.CLEAN)
    store.upsert(document, b"%PDF-1.4", [])
    store.correct(document.document_id, "part_i.city", "Nashville", "Ada Reviewer")
    assert verify_database(db) == []
    conn = sqlite3.connect(db)
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("DELETE FROM audit_log")
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("UPDATE audit_log SET actor = 'forged'")
    conn.close()


def test_formula_injection_is_quoted() -> None:
    assert excel_safe("=1+1") == "'=1+1"
    assert excel_safe("+cmd") == "'+cmd"
    assert excel_safe("-2+3") == "'-2+3"
    assert excel_safe("-5.00") == "-5.00"
    assert excel_safe("@cmd") == "'@cmd"
    assert excel_safe("\tcmd") == "'\tcmd"
    assert excel_safe("\rcmd") == "'\rcmd"
    document = build_document(random.Random(6), 1, Variant.CLEAN)
    payload = document.model_dump(mode="json")
    payload["part_i"]["name"]["value"] = '=HYPERLINK("http://example.test")'
    loaded = ScheduleK1.model_validate(payload)
    record = ExportRecord(document=loaded, approved_by="Ada Reviewer", approved_at="2026-01-01")
    csv_text = records_to_csv([record])
    assert "'=HYPERLINK" in csv_text
    for marker in ("@sum", "\t=1", "\r=1"):
        payload["part_ii"]["name"]["value"] = marker
        quoted = ScheduleK1.model_validate(payload)
        text = records_to_csv(
            [ExportRecord(document=quoted, approved_by="Ada Reviewer", approved_at="2026-01-01")]
        )
        assert f"'{marker}" in text
    sheet = load_workbook(BytesIO(records_to_xlsx([record])))["approved_k1s"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert any(isinstance(value, str) and value.startswith("'=HYPERLINK") for value in values)


def test_unrecognized_cap_follows_the_threshold() -> None:
    assert below_review_threshold(0.8) == 0.49
    assert below_review_threshold(0.4) < 0.4
    assert below_review_threshold(0.4) == pytest.approx(0.39)
    observed = {
        "VARIANT": [Observed("memo", 1, None, "deterministic")],
        "5 Interest income": [Observed("10.00", 1, None, "deterministic", label_score=1.0)],
    }
    document = assemble("K1-NOTE", observed, extractor="deterministic", threshold=0.4)
    assert document.layout_recognized is False
    assert document.part_iii.interest_income.confidence < 0.4


def test_an_unrecognized_layout_keeps_a_clean_amount_under_review() -> None:
    from k1lens.extract.labels import LayoutStats

    document = assemble(
        "K1-CAP",
        {
            "1 Ordinary business income": [
                Observed("15.00", 1, None, "deterministic", label_score=1.0, proximity=1.0)
            ]
        },
        extractor="deterministic",
        layout_stats=LayoutStats(seen_lines=12, strong_anchors=1),
    )
    assert document.layout_recognized is False
    assert document.part_iii.ordinary_business_income.value == Decimal("15.00")
    assert document.part_iii.ordinary_business_income.confidence < REVIEW_THRESHOLD


def test_preview_png_is_not_a_black_page(tmp_path: Path) -> None:
    pdf = tmp_path / "letter.pdf"
    _write_pdf(pdf, ["This is a synthetic letter.", "It is not a Schedule K-1."])
    image = Image.open(BytesIO(pdf_preview_png(pdf.read_bytes())))
    pixels = list(image.get_flattened_data())
    white = sum(1 for pixel in pixels if pixel[0] > 240 and pixel[1] > 240 and pixel[2] > 240)
    black = sum(1 for pixel in pixels if pixel == (0, 0, 0))
    assert white > black
    assert white / len(pixels) > 0.5
